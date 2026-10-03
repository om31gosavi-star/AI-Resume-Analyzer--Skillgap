"""SkillGap - find the skills missing from your resume for a target job."""
import io
import json
import os

from flask import Flask, jsonify, request, send_from_directory

import db
from analyzer import analyze, build_matchers, extract_skills

FRONTEND_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "frontend")
MAX_UPLOAD_MB = 5

app = Flask(__name__, static_folder=None)
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_MB * 1024 * 1024
db.init_db()


def _skill_dict(row, importance):
    return {"id": row["id"], "name": row["name"], "category": row["category"],
            "resource": row["resource"], "importance": importance}


# ---------- Frontend ----------
@app.get("/")
def index():
    return send_from_directory(FRONTEND_DIR, "index.html")


@app.get("/<path:filename>")
def static_files(filename):
    return send_from_directory(FRONTEND_DIR, filename)


# ---------- API ----------
@app.get("/api/roles")
def roles():
    conn = db.get_conn()
    rows = conn.execute(
        """SELECT r.id, r.title, r.description, COUNT(js.skill_id) AS skill_count
           FROM job_roles r LEFT JOIN job_skills js ON js.role_id = r.id
           GROUP BY r.id ORDER BY r.title""").fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])


@app.get("/api/skills")
def skills():
    conn = db.get_conn()
    rows = conn.execute("SELECT id, name, category FROM skills ORDER BY category, name").fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])


@app.post("/api/extract-text")
def extract_text():
    """Upload a resume (.pdf, .docx, .txt) and get plain text back."""
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify(error="No file uploaded"), 400
    name = f.filename.lower()
    data = f.read()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            text = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
        elif name.endswith(".docx"):
            import docx
            d = docx.Document(io.BytesIO(data))
            text = "\n".join(p.text for p in d.paragraphs)
            for t in d.tables:
                for row in t.rows:
                    text += "\n" + " ".join(c.text for c in row.cells)
        elif name.endswith(".txt"):
            text = data.decode("utf-8", errors="ignore")
        else:
            return jsonify(error="Unsupported file type. Use PDF, DOCX or TXT."), 400
    except Exception as e:  # corrupted file etc.
        return jsonify(error=f"Could not read file: {e}"), 400
    if not text.strip():
        return jsonify(error="No text found in the file (scanned PDFs are not supported)."), 400
    return jsonify(text=text)


@app.post("/api/analyze")
def analyze_resume():
    body = request.get_json(silent=True) or {}
    resume = (body.get("resume_text") or "").strip()
    role_id = body.get("role_id")
    jd = (body.get("job_description") or "").strip()

    if len(resume) < 20:
        return jsonify(error="Please paste or upload your resume text."), 400
    if not role_id and not jd:
        return jsonify(error="Choose a job role or paste a job description."), 400

    conn = db.get_conn()
    skill_rows = conn.execute("SELECT * FROM skills").fetchall()
    matchers = build_matchers(skill_rows)

    if jd:
        target = "Custom job description"
        by_id = {r["id"]: r for r in skill_rows}
        ids = extract_skills(jd, matchers)
        if not ids:
            conn.close()
            return jsonify(error="No known skills were found in that job description."), 400
        # skills the JD mentions are treated as required
        target_skills = [_skill_dict(by_id[i], "required") for i in ids]
    else:
        role = conn.execute("SELECT * FROM job_roles WHERE id = ?", (role_id,)).fetchone()
        if not role:
            conn.close()
            return jsonify(error="Unknown job role."), 404
        target = role["title"]
        rows = conn.execute(
            """SELECT s.*, js.importance FROM job_skills js
               JOIN skills s ON s.id = js.skill_id WHERE js.role_id = ?""", (role_id,)).fetchall()
        target_skills = [_skill_dict(r, r["importance"]) for r in rows]

    result = analyze(resume, target_skills, matchers)
    result["target"] = target

    cur = conn.execute(
        "INSERT INTO analyses (target, match_score, matched_json, missing_json) VALUES (?,?,?,?)",
        (target, result["score"],
         json.dumps([s["name"] for s in result["matched"]]),
         json.dumps([s["name"] for s in result["missing"]])))
    conn.commit()
    result["id"] = cur.lastrowid
    conn.close()
    return jsonify(result)


@app.get("/api/history")
def history():
    conn = db.get_conn()
    rows = conn.execute(
        "SELECT * FROM analyses ORDER BY id DESC LIMIT 20").fetchall()
    conn.close()
    out = []
    for r in rows:
        out.append({"id": r["id"], "created_at": r["created_at"], "target": r["target"],
                    "score": r["match_score"],
                    "matched": json.loads(r["matched_json"]),
                    "missing": json.loads(r["missing_json"])})
    return jsonify(out)


@app.delete("/api/history/<int:analysis_id>")
def delete_history(analysis_id):
    conn = db.get_conn()
    conn.execute("DELETE FROM analyses WHERE id = ?", (analysis_id,))
    conn.commit()
    conn.close()
    return jsonify(ok=True)


@app.errorhandler(413)
def too_large(_):
    return jsonify(error=f"File too large (max {MAX_UPLOAD_MB} MB)."), 413


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    print(f"SkillGap running at http://localhost:{port}")
    app.run(host="127.0.0.1", port=port, debug=True)
